"""Verify exact report values, portable artifacts and receipt boundaries."""

from __future__ import annotations

import hashlib
import json
import zipfile
from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from expense_claim_helper.core.reporting import ReportError, build_reports


@pytest.fixture
def claim(tmp_path: Path) -> dict:
    first = tmp_path / "fictional lodging.png"
    second = tmp_path / "fictional train.pdf"
    first.write_bytes(b"fictional receipt image bytes")
    second.write_bytes(b"fictional receipt PDF bytes")
    return {
        "code": "EX-2026-001",
        "created_at": "2026-10-15T10:00:00+08:00",
        "status": "draft",
        "currency": "CNY",
        "policy_sha": "a" * 64,
        "policy_yaml": "name: Fictional travel standard\nmeal_allowance_cents: 5000\n",
        "trip_ids": [1],
        "details": [
            {
                "sequence": 1,
                "trip_id": 1,
                "date": "2026-10-13",
                "destination": "City A",
                "category": "lodging",
                "description": "Fictional lodging",
                "amount_cents": 38001,
                "eligible_cents": 35000,
                "self_paid_cents": 3001,
                "receipt_path": str(first),
                "flags": ["Over lodging limit"],
            },
            {
                "sequence": 2,
                "trip_id": 1,
                "date": "2026-10-13",
                "destination": "City A",
                "category": "train",
                "description": "Fictional train fare",
                "amount_cents": 12345,
                "eligible_cents": 12345,
                "self_paid_cents": 0,
                "receipt_path": str(second),
                "flags": [],
            },
            {
                "sequence": 3,
                "trip_id": 1,
                "date": "2026-10-13",
                "destination": "City A",
                "category": "meal_allowance",
                "description": "One fictional travel day",
                "amount_cents": 0,
                "eligible_cents": 5000,
                "self_paid_cents": 0,
                "receipt_path": "",
                "flags": [],
            },
        ],
        "totals": {
            "amount_cents": 50346,
            "eligible_cents": 52345,
            "self_paid_cents": 3001,
            "allowance_cents": 5000,
            "advance_cents": 50346,
        },
    }


def test_all_formats_exact_totals_and_manifest(claim: dict, tmp_path: Path) -> None:
    original = deepcopy(claim)
    outputs = build_reports(claim, tmp_path / "reports", ["xlsx", "pdf", "docx", "zip"])
    assert claim == original
    assert set(outputs) == {"xlsx", "pdf", "docx", "zip"}
    assert all(
        Path(path).is_absolute() and Path(path).stat().st_size > 0 for path in outputs.values()
    )
    book = load_workbook(outputs["xlsx"], data_only=False)
    assert book["Summary"]["B11"].value == 503.46
    assert book["Summary"]["C11"].value == 50346
    assert book["Summary"]["B13"].value == 523.45
    assert book["Details"]["G5"].value == 503.46
    assert book["Details"]["H5"].value == 523.45
    assert book["Details"]["I5"].value == 30.01
    assert book["Details"]["L5"].value == 50346
    assert book["Details"].print_title_rows == "$1:$1"
    assert "Policy snapshot" in book.sheetnames
    assert book["Details"]["G2"].font.color.rgb == "00B63A3A"
    pdf = PdfReader(outputs["pdf"])
    text = "\n".join(page.extract_text() for page in pdf.pages)
    for expected in ["503.46", "523.45", "30.01", "Prepared by", "Policy snapshot"]:
        assert expected in text
    assert str(tmp_path) not in text
    document = Document(outputs["docx"])
    word_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Prepared by:" in word_text
    assert document.tables[0].rows[3].cells[1].text == "523.45"
    assert "w:tblHeader" in document.tables[1].rows[0]._tr.xml
    with zipfile.ZipFile(outputs["zip"]) as archive:
        names = archive.namelist()
        assert [name for name in names if name.startswith("receipts/")] == [
            "receipts/receipt_001_fictional_lodging.png",
            "receipts/receipt_002_fictional_train.pdf",
        ]
        manifest = json.loads(archive.read("manifest.json"))
        assert manifest["totals"] == claim["totals"]
        for receipt in manifest["receipts"]:
            data = archive.read(f"receipts/{receipt['filename']}")
            assert hashlib.sha256(data).hexdigest() == receipt["sha256"]
        for report in manifest["reports"]:
            assert hashlib.sha256(archive.read(report["filename"])).hexdigest() == report["sha256"]
        assert str(tmp_path) not in archive.read("manifest.json").decode()


def test_zip_alone_embeds_both_documents(claim: dict, tmp_path: Path) -> None:
    outputs = build_reports(claim, tmp_path / "zip-only", ["zip"])
    assert list(outputs) == ["zip"]
    assert len(list((tmp_path / "zip-only").iterdir())) == 1
    with zipfile.ZipFile(outputs["zip"]) as archive:
        assert f"claim_{claim['code']}.pdf" in archive.namelist()
        assert f"claim_{claim['code']}.xlsx" in archive.namelist()


def test_duplicate_receipt_path_has_one_number(claim: dict, tmp_path: Path) -> None:
    claim["details"][1]["receipt_path"] = claim["details"][0]["receipt_path"]
    output = build_reports(claim, tmp_path / "unique", ["zip"])
    with zipfile.ZipFile(output["zip"]) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        assert len(manifest["receipts"]) == 1
        assert manifest["receipts"][0]["detail_sequences"] == [1, 2]


def test_formula_like_text_is_not_a_formula(claim: dict, tmp_path: Path) -> None:
    claim["details"][0]["description"] = '=HYPERLINK("https://invalid.example", "label")'
    claim["details"][0]["destination"] = "+SUM(1,2)"
    claim["details"][1]["flags"] = ["@SUM(1,2)"]
    output = build_reports(claim, tmp_path / "strings", ["xlsx"])
    book = load_workbook(output["xlsx"], data_only=False)
    assert book["Details"]["F2"].value.startswith("=")
    for sheet in book:
        for row in sheet:
            assert all(cell.data_type != "f" for cell in row)


@pytest.mark.parametrize("change", ["total", "float", "negative", "boolean", "sequence", "code"])
def test_invalid_claim_rejected(claim: dict, tmp_path: Path, change: str) -> None:
    if change == "total":
        claim["totals"]["amount_cents"] += 1
    elif change == "float":
        claim["details"][0]["amount_cents"] = 380.01
    elif change == "negative":
        claim["details"][0]["self_paid_cents"] = -1
    elif change == "boolean":
        claim["totals"]["advance_cents"] = True
    elif change == "sequence":
        claim["details"][1]["sequence"] = 1
    else:
        claim["code"] = "../escape"
    with pytest.raises(ReportError):
        build_reports(claim, tmp_path / "invalid", ["xlsx"])
    assert not (tmp_path / "invalid").exists()


def test_missing_and_symlink_receipts_are_rejected(claim: dict, tmp_path: Path) -> None:
    source = Path(claim["details"][0]["receipt_path"])
    link = tmp_path / "link.png"
    link.symlink_to(source)
    claim["details"][0]["receipt_path"] = str(link)
    with pytest.raises(ReportError, match="without symlinks"):
        build_reports(claim, tmp_path / "symlink", ["zip"])
    claim["details"][0]["receipt_path"] = str(tmp_path / "missing.png")
    with pytest.raises(ReportError, match="must exist"):
        build_reports(claim, tmp_path / "missing", ["pdf"])
    folder = tmp_path / "linked-directory"
    folder.symlink_to(tmp_path, target_is_directory=True)
    claim["details"][0]["receipt_path"] = str(folder / source.name)
    with pytest.raises(ReportError, match="without symlinks"):
        build_reports(claim, tmp_path / "symlink-dir", ["zip"])


def test_outputs_never_overwrite(claim: dict, tmp_path: Path) -> None:
    folder = tmp_path / "existing"
    outputs = build_reports(claim, folder, ["xlsx"])
    original = Path(outputs["xlsx"]).read_bytes()
    with pytest.raises(ReportError, match="already exists"):
        build_reports(claim, folder, ["xlsx"])
    assert Path(outputs["xlsx"]).read_bytes() == original
    folder_link = tmp_path / "linked-output"
    folder_link.symlink_to(folder, target_is_directory=True)
    with pytest.raises(ReportError, match="symlinks"):
        build_reports(claim, folder_link, ["pdf"])


def test_paged_pdf_repeats_header_and_preserves_last_row(claim: dict, tmp_path: Path) -> None:
    template = claim["details"][0]
    claim["details"] = []
    for number in range(1, 66):
        row = {
            **template,
            "sequence": number,
            "description": f"Fictional expense {number:03d}. " + "A long wrapped description. " * 4,
        }
        claim["details"].append(row)
    claim["totals"] = {
        key: sum(row[key] for row in claim["details"])
        for key in ("amount_cents", "eligible_cents", "self_paid_cents")
    }
    outputs = build_reports(claim, tmp_path / "pages", ["pdf"])
    reader = PdfReader(outputs["pdf"])
    pages = [page.extract_text() for page in reader.pages]
    assert len(pages) >= 5
    assert "Fictional expense 065" in "\n".join(pages)
    detail_pages = [page for page in pages if "Fictional expense" in page]
    assert all("Recorded" in page and "Self-paid" in page for page in detail_pages)
    assert all(f"Page {number}" in text for number, text in enumerate(pages, 1))


def test_unsupported_format_is_clear(claim: dict, tmp_path: Path) -> None:
    with pytest.raises(ReportError, match="Formats"):
        build_reports(claim, tmp_path, ["html"])
