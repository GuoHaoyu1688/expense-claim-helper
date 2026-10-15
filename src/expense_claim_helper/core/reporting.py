"""Create local, auditable claim documents from a frozen claim snapshot."""

from __future__ import annotations

import hashlib
import json
import os
import re
import stat
import tempfile
import zipfile
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

FORMATS = frozenset({"xlsx", "pdf", "docx", "zip"})
MONEY_KEYS = ("amount_cents", "eligible_cents", "self_paid_cents")
MAX_RECEIPT_BYTES = 100 * 1024 * 1024


class ReportError(ValueError):
    """A claim or receipt cannot be safely exported."""


@dataclass(frozen=True)
class Receipt:
    number: int
    name: str
    original_name: str
    sequences: tuple[int, ...]
    data: bytes

    def manifest(self) -> dict[str, Any]:
        return {
            "number": self.number,
            "filename": self.name,
            "original_filename": self.original_name,
            "detail_sequences": list(self.sequences),
            "size_bytes": len(self.data),
            "sha256": hashlib.sha256(self.data).hexdigest(),
        }


def _text(value: Any, maximum: int = 4000) -> str:
    text = "" if value is None else str(value)
    # XML used by all three document formats excludes these control characters.
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    if len(text) > maximum:
        raise ReportError(f"A report field exceeds {maximum} characters.")
    return text


def _cents(value: Any, name: str) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        raise ReportError(f"{name} must be a nonnegative integer number of cents.")
    if value > 999_999_999_999_999:
        raise ReportError(f"{name} exceeds the exact spreadsheet integer limit.")
    return value


def _money(value: int) -> Decimal:
    return Decimal(value) / Decimal(100)


def _money_text(value: int) -> str:
    return f"{_money(value):,.2f}"


def _safe_read(path: Path) -> bytes:
    """Open each POSIX path component without following a symlink."""
    absolute = Path(os.path.abspath(path))
    if not hasattr(os, "O_NOFOLLOW"):
        raise ReportError("Safe receipt reading requires a platform with O_NOFOLLOW.")
    directory_fd = os.open(absolute.anchor, os.O_RDONLY | os.O_DIRECTORY)
    file_fd: int | None = None
    try:
        for component in absolute.parts[1:-1]:
            next_fd = os.open(
                component,
                os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW,
                dir_fd=directory_fd,
            )
            os.close(directory_fd)
            directory_fd = next_fd
        file_fd = os.open(
            absolute.name, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK, dir_fd=directory_fd
        )
        metadata = os.fstat(file_fd)
        if not stat.S_ISREG(metadata.st_mode):
            raise ReportError(f"Receipt is not a regular file: {path.name}")
        if metadata.st_size > MAX_RECEIPT_BYTES:
            raise ReportError(f"Receipt exceeds the 100 MiB limit: {path.name}")
        with os.fdopen(file_fd, "rb", closefd=False) as handle:
            data = handle.read(MAX_RECEIPT_BYTES + 1)
        if len(data) > MAX_RECEIPT_BYTES:
            raise ReportError(f"Receipt exceeds the 100 MiB limit: {path.name}")
        return data
    except OSError as exc:
        raise ReportError(
            f"Cannot read receipt {path.name!r}; it must exist without symlinks."
        ) from exc
    finally:
        if file_fd is not None:
            os.close(file_fd)
        os.close(directory_fd)


def _prepare(claim: dict[str, Any]) -> dict[str, Any]:
    code = claim.get("code", "")
    if not isinstance(code, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,63}", code):
        raise ReportError("Claim code must be 1-64 safe ASCII letters, digits, dots or hyphens.")
    raw_rows = claim.get("details", claim.get("rows"))
    if not isinstance(raw_rows, list) or not raw_rows:
        raise ReportError("The claim must contain at least one detail row.")
    rows: list[dict[str, Any]] = []
    seen_sequences: set[int] = set()
    for position, raw in enumerate(raw_rows, 1):
        row = dict(raw)
        sequence = row.get("sequence", position)
        if isinstance(sequence, bool) or not isinstance(sequence, int) or sequence <= 0:
            raise ReportError("Each detail sequence must be a positive integer.")
        if sequence in seen_sequences:
            raise ReportError("Detail sequences must be unique.")
        seen_sequences.add(sequence)
        row["sequence"] = sequence
        for key in MONEY_KEYS:
            row[key] = _cents(row.get(key, 0), key)
        for key in ("trip_id", "date", "destination", "category", "description"):
            row[key] = _text(row.get(key, ""))
        flags = row.get("flags", [])
        if not isinstance(flags, list):
            raise ReportError("Detail flags must be a list of strings.")
        row["flags"] = [_text(flag) for flag in flags]
        row["receipt_label"] = "No receipt" if row.get("amount_cents") else "Not applicable"
        rows.append(row)
    rows.sort(key=lambda row: row["sequence"])
    totals = dict(claim.get("totals", {}))
    for key in MONEY_KEYS:
        expected = sum(row[key] for row in rows)
        totals[key] = _cents(totals.get(key, expected), key)
        if totals[key] != expected:
            raise ReportError(f"Claim total {key} does not equal its detail rows.")
    for key in ("advance_cents", "allowance_cents"):
        totals[key] = _cents(totals.get(key, 0), key)
    receipts: list[Receipt] = []
    path_rows: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        if row.get("receipt_path"):
            absolute = os.path.abspath(str(row["receipt_path"]))
            path_rows.setdefault(absolute, []).append(row)
    for number, (path_text, matching_rows) in enumerate(path_rows.items(), 1):
        path = Path(path_text)
        original_name = _text(path.name, 255)
        safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", path.name).strip("._")[:120] or "receipt"
        name = f"receipt_{number:03d}_{safe_name}"
        receipts.append(
            Receipt(
                number,
                name,
                original_name,
                tuple(row["sequence"] for row in matching_rows),
                _safe_read(path),
            )
        )
        for row in matching_rows:
            row["receipt_label"] = f"{number:03d}: {original_name}"
    return {
        **claim,
        "code": code,
        "created_at": _text(claim.get("created_at", "")),
        "status": _text(claim.get("status", "draft")),
        "currency": _text(claim.get("currency", "CNY"), 12),
        "policy_sha": _text(claim.get("policy_sha", ""), 128),
        "policy_text": _text(claim.get("policy_yaml", claim.get("policy_text", "")), 100000),
        "details": rows,
        "totals": totals,
        "receipts": receipts,
    }


def _summary_rows(claim: dict[str, Any]) -> list[tuple[str, int]]:
    totals = claim["totals"]
    return [
        ("Recorded expenses", totals["amount_cents"]),
        ("Allowance entitlement", totals["allowance_cents"]),
        ("Eligible reimbursement", totals["eligible_cents"]),
        ("Employee self-paid difference", totals["self_paid_cents"]),
        ("Employee advance paid", totals["advance_cents"]),
    ]


def _write_xlsx(claim: dict[str, Any], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.pagebreak import Break

    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    navy = "163D37"
    pale = "E9F0EB"
    red = "B63A3A"

    def cell(sheet: Any, row: int, column: int, value: Any) -> None:
        item = sheet.cell(row, column, value)
        if isinstance(value, str):
            # Explicit string storage prevents formula injection without altering visible text.
            item.data_type = "s"
        item.alignment = Alignment(vertical="top", wrap_text=True)

    def header(sheet: Any, row: int, labels: list[str]) -> None:
        for index, label in enumerate(labels, 1):
            cell(sheet, row, index, label)
            item = sheet.cell(row, index)
            item.font = Font(color="FFFFFF", bold=True)
            item.fill = PatternFill("solid", fgColor=navy)
        sheet.row_dimensions[row].height = 28

    cell(summary, 1, 1, "TRAVEL EXPENSE CLAIM")
    summary["A1"].font = Font(size=20, bold=True, color=navy)
    summary.merge_cells("A1:D1")
    for row, (label, value) in enumerate(
        [
            ("Claim code", claim["code"]),
            ("Created", claim["created_at"]),
            ("Status", claim["status"]),
            ("Currency", claim["currency"]),
            ("Trips", ", ".join(str(value) for value in claim.get("trip_ids", []))),
            ("Policy SHA-256", claim["policy_sha"]),
        ],
        3,
    ):
        cell(summary, row, 1, label)
        cell(summary, row, 2, value)
        summary.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    header(summary, 10, ["Total", f"Amount ({claim['currency']})", "Exact cents"])
    for row, (label, cents) in enumerate(_summary_rows(claim), 11):
        cell(summary, row, 1, label)
        cell(summary, row, 2, _money(cents))
        summary.cell(row, 2).number_format = "#,##0.00;[Red]-#,##0.00"
        cell(summary, row, 3, cents)
    cell(summary, 17, 1, "Prepared by: ____________________")
    cell(summary, 19, 1, "Approved by: ____________________")
    cell(summary, 21, 1, "Finance review: ____________________")
    cell(summary, 23, 1, "Payment date: ____________________")
    cell(summary, 25, 1, "Allowances are entitlements, not additional recorded cash expenses.")
    summary.merge_cells("A25:D26")
    for column, width in {"A": 36, "B": 25, "C": 23, "D": 12}.items():
        summary.column_dimensions[column].width = width
    summary.print_options.horizontalCentered = True
    summary.print_area = "A1:D26"

    details = book.create_sheet("Details")
    labels = [
        "No.",
        "Trip",
        "Date",
        "Destination",
        "Category",
        "Description",
        "Recorded",
        "Eligible",
        "Self-paid",
        "Receipt",
        "Flags",
        "Recorded cents",
        "Eligible cents",
        "Self-paid cents",
    ]
    header(details, 1, labels)
    for row_number, row in enumerate(claim["details"], 2):
        values = [
            row["sequence"],
            row["trip_id"],
            row["date"],
            row["destination"],
            row["category"],
            row["description"],
            *[_money(row[key]) for key in MONEY_KEYS],
            row["receipt_label"],
            "; ".join(row["flags"]),
            *[row[key] for key in MONEY_KEYS],
        ]
        for column, value in enumerate(values, 1):
            cell(details, row_number, column, value)
            if column in (7, 8, 9):
                details.cell(row_number, column).number_format = "#,##0.00;[Red]-#,##0.00"
            if row["self_paid_cents"] or row["flags"]:
                details.cell(row_number, column).font = Font(color=red)
            elif row_number % 2 == 0:
                details.cell(row_number, column).fill = PatternFill("solid", fgColor=pale)
        details.row_dimensions[row_number].height = min(
            300, max(36, 15 * (1 + max(len(str(value)) // 30 for value in values)))
        )
    total_row = len(claim["details"]) + 2
    cell(details, total_row, 6, "TOTAL")
    for column, key in zip((7, 8, 9), MONEY_KEYS, strict=True):
        cell(details, total_row, column, _money(claim["totals"][key]))
        details.cell(total_row, column).number_format = "#,##0.00"
        details.cell(total_row, column).font = Font(bold=True, color=navy)
        cell(details, total_row, column + 5, claim["totals"][key])
    widths = [7, 10, 13, 19, 17, 42, 16, 16, 16, 29, 29, 20, 20, 20]
    for number, width in enumerate(widths, 1):
        details.column_dimensions[get_column_letter(number)].width = width
        if number >= 12:
            details.column_dimensions[get_column_letter(number)].hidden = True
    details.freeze_panes = "G2"
    details.auto_filter.ref = f"A1:N{total_row - 1}"
    details.print_title_rows = "1:1"
    details.print_area = f"A1:K{total_row}"
    details.page_setup.orientation = "landscape"
    details.page_setup.paperSize = details.PAPERSIZE_A3
    details.page_setup.fitToWidth = 1
    details.page_setup.fitToHeight = 0
    details.sheet_properties.pageSetUpPr.fitToPage = True
    for number in range(29, total_row, 28):
        details.row_breaks.append(Break(id=number))

    receipts = book.create_sheet("Receipts")
    header(
        receipts,
        1,
        ["Receipt no.", "Export filename", "Original filename", "Detail rows", "SHA-256"],
    )
    for row_number, receipt in enumerate(claim["receipts"], 2):
        manifest = receipt.manifest()
        for column, value in enumerate(
            [
                receipt.number,
                receipt.name,
                receipt.original_name,
                ", ".join(map(str, receipt.sequences)),
                manifest["sha256"],
            ],
            1,
        ):
            cell(receipts, row_number, column, value)
    for column, width in {"A": 16, "B": 45, "C": 40, "D": 20, "E": 68}.items():
        receipts.column_dimensions[column].width = width
    receipts.freeze_panes = "A2"
    if claim["policy_text"]:
        policy = book.create_sheet("Policy snapshot")
        cell(policy, 1, 1, "Policy SHA-256")
        cell(policy, 1, 2, claim["policy_sha"])
        for row_number, line in enumerate(claim["policy_text"].splitlines(), 3):
            cell(policy, row_number, 1, line)
        policy.column_dimensions["A"].width = 100
        policy.column_dimensions["B"].width = 68
    for sheet in book:
        sheet.sheet_view.showGridLines = False
        sheet.oddFooter.center.text = f"Claim {claim['code']} | Page &P of &N"
    book.save(path)


def _write_pdf(claim: dict[str, Any], path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (
        KeepTogether,
        LongTable,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    # ReportLab's bundled CID metrics support fictional Chinese city/category text offline.
    font = "STSong-Light"
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(font))
    ink = colors.HexColor("#163D37")
    pale = colors.HexColor("#F0F4F1")
    red = colors.HexColor("#AA3333")
    styles = getSampleStyleSheet()
    body = ParagraphStyle("ClaimBody", fontName="Helvetica", fontSize=9, leading=13, textColor=ink)
    small = ParagraphStyle("ClaimSmall", parent=body, fontSize=7.3, leading=10, wordWrap="CJK")
    right = ParagraphStyle("ClaimMoney", parent=small, alignment=TA_RIGHT)
    white = ParagraphStyle("ClaimHead", parent=small, textColor=colors.white)
    title = ParagraphStyle(
        "ClaimTitle", parent=styles["Title"], fontSize=23, textColor=ink, alignment=0, spaceAfter=12
    )
    section = ParagraphStyle(
        "ClaimSection", parent=styles["Heading2"], textColor=ink, spaceBefore=12, spaceAfter=8
    )

    def p(value: Any, style: Any = small) -> Any:
        if re.search(r"[\u2e80-\u9fff]", str(value)):
            style = ParagraphStyle("CJK", parent=style, fontName=font)
        return Paragraph(escape(str(value)).replace("\n", "<br/>"), style)

    document = SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        leftMargin=34,
        rightMargin=34,
        topMargin=32,
        bottomMargin=34,
        title=f"Expense claim {claim['code']}",
        author="Expense Claim Helper",
    )
    story: list[Any] = [
        p("TRAVEL EXPENSE CLAIM", title),
        p(f"{claim['code']}  |  {claim['status']}  |  {claim['currency']}", body),
        p(
            f"Created: {claim['created_at']}  |  Trips: "
            + ", ".join(str(value) for value in claim.get("trip_ids", [])),
            body,
        ),
        Spacer(1, 13),
    ]
    figure = ParagraphStyle("ClaimFigure", parent=body, fontSize=17, leading=22)
    totals = [
        [p(label, body) for label, _ in _summary_rows(claim)],
        [p(_money_text(cents), figure) for _, cents in _summary_rows(claim)],
    ]
    total_table = Table(totals, colWidths=[document.width / 5] * 5)
    total_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), pale),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    total_table.hAlign = "LEFT"
    story.extend(
        [
            total_table,
            Spacer(1, 8),
            p("Allowances are entitlements, not additional recorded cash expenses.", body),
            p("Expense detail", section),
        ]
    )
    labels = [
        "No.",
        "Trip / date",
        "Destination / category",
        "Description",
        "Recorded",
        "Eligible",
        "Self-paid",
        "Receipt / flags",
    ]
    table_rows = [[p(label, white) for label in labels]]
    widths = [28, 66, 94, 190, 66, 66, 66, document.width - 576]
    for row in claim["details"]:
        description_style = (
            ParagraphStyle("Flagged", parent=small, textColor=red)
            if row["self_paid_cents"] or row["flags"]
            else small
        )
        table_rows.append(
            [
                p(row["sequence"]),
                p(f"{row['trip_id']}\n{row['date']}"),
                p(f"{row['destination']}\n{row['category']}"),
                p(row["description"], description_style),
                *[p(_money_text(row[key]), right) for key in MONEY_KEYS],
                p(
                    row["receipt_label"] + ("\n" + "; ".join(row["flags"]) if row["flags"] else ""),
                    description_style,
                ),
            ]
        )
    table_rows.append(
        [
            p(""),
            p(""),
            p(""),
            p("TOTAL"),
            *[p(_money_text(claim["totals"][key]), right) for key in MONEY_KEYS],
            p(""),
        ]
    )
    detail_table = LongTable(table_rows, colWidths=widths, repeatRows=1, splitInRow=1)
    detail_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), ink),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, pale]),
                ("LINEBELOW", (0, 0), (-1, 0), 0.5, ink),
                ("LINEABOVE", (0, -1), (-1, -1), 0.7, ink),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.append(detail_table)
    story.append(p("Receipt register", section))
    receipt_rows = [
        [
            p(label, white)
            for label in ["No.", "Export filename", "Original filename", "Detail rows"]
        ]
    ]
    for receipt in claim["receipts"]:
        receipt_rows.append(
            [
                p(receipt.number),
                p(receipt.name),
                p(receipt.original_name),
                p(", ".join(map(str, receipt.sequences))),
            ]
        )
    if not claim["receipts"]:
        receipt_rows.append([p("-"), p("No receipt files attached"), p(""), p("")])
    receipt_table = LongTable(
        receipt_rows, colWidths=[35, 300, document.width - 410, 75], repeatRows=1, splitInRow=1
    )
    receipt_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), ink),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, pale]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.extend(
        [
            receipt_table,
            Spacer(1, 8),
            KeepTogether(
                [
                    p("Approval", body),
                    Spacer(1, 6),
                    p(
                        "Prepared by: _______________    Approved by: _______________    "
                        "Finance review: _______________    Payment date: _______________",
                        body,
                    ),
                ]
            ),
            Spacer(1, 8),
            p(f"Policy SHA-256: {claim['policy_sha']}"),
        ]
    )
    if claim["policy_text"]:
        story.extend([PageBreak(), p("Policy snapshot", title)])
        for line in claim["policy_text"].splitlines():
            story.append(p(line or " ", small))

    def footer(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(ink)
        canvas.drawString(34, 18, f"Claim {claim['code']}")
        canvas.drawRightString(landscape(A4)[0] - 34, 18, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)


def _write_docx(claim: dict[str, Any], path: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = section.bottom_margin = Inches(0.5)
    section.left_margin = section.right_margin = Inches(0.5)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor.from_string("163D37")
    document.add_heading("Travel expense claim", 0)
    document.add_paragraph(f"{claim['code']} | {claim['status']} | {claim['currency']}")
    document.add_paragraph(f"Created: {claim['created_at']}")
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Total"
    table.rows[0].cells[1].text = f"Amount ({claim['currency']})"
    for label, cents in _summary_rows(claim):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, _money_text(cents)
    document.add_paragraph("Allowances are entitlements, not additional recorded cash expenses.")
    document.add_heading("Expense detail", level=1)
    labels = [
        "No.",
        "Trip / date",
        "Destination / category",
        "Description",
        "Recorded",
        "Eligible",
        "Self-paid",
        "Receipt / flags",
    ]
    table = document.add_table(rows=1, cols=len(labels))
    table.style = "Light Shading Accent 1"
    for cell, label in zip(table.rows[0].cells, labels, strict=True):
        cell.text = label
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    widths = [0.4, 0.85, 1.15, 2.8, 0.9, 0.9, 0.9, 1.8]
    table.autofit = False
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Inches(width)
    for row in claim["details"]:
        cells = table.add_row().cells
        values = [
            str(row["sequence"]),
            f"{row['trip_id']}\n{row['date']}",
            f"{row['destination']}\n{row['category']}",
            row["description"],
            *[_money_text(row[key]) for key in MONEY_KEYS],
            row["receipt_label"] + "\n" + "; ".join(row["flags"]),
        ]
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
            if row["self_paid_cents"] or row["flags"]:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor.from_string("B63A3A")
    cells = table.add_row().cells
    cells[3].text = "TOTAL"
    for index, key in enumerate(MONEY_KEYS, 4):
        cells[index].text = _money_text(claim["totals"][key])
    document.add_heading("Receipt register", level=1)
    for receipt in claim["receipts"]:
        document.add_paragraph(
            f"{receipt.number:03d}: {receipt.name} | Original: {receipt.original_name} | "
            f"Detail rows: {', '.join(map(str, receipt.sequences))}"
        )
    if not claim["receipts"]:
        document.add_paragraph("No receipt files attached.")
    document.add_heading("Approval", level=1)
    document.add_paragraph("Prepared by: ____________________   Approved by: ____________________")
    document.add_paragraph(
        "Finance review: ____________________   Payment date: ____________________"
    )
    document.add_paragraph(f"Policy SHA-256: {claim['policy_sha']}")
    if claim["policy_text"]:
        document.add_page_break()
        document.add_heading("Policy snapshot", level=1)
        document.add_paragraph(claim["policy_text"])
    footer = section.footer.paragraphs[0]
    footer.text = f"Claim {claim['code']} | Page "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = f"Expense claim {claim['code']}"
    document.core_properties.author = "Expense Claim Helper"
    document.save(path)


def build_reports(claim: dict[str, Any], out_dir: Path, formats: list[str]) -> dict[str, str]:
    """Write requested reports, returning their absolute paths.

    ZIP always embeds PDF, XLSX, a hash manifest and numbered receipt copies.
    DOCX is embedded when requested. PDF/XLSX generated for ZIP alone are only
    inside the ZIP. All documents are generated before publishing final outputs.
    Existing symlinks are rejected, as are symlinks anywhere in receipt paths.
    """
    if not formats or any(item not in FORMATS for item in formats):
        raise ReportError("Formats must be a nonempty list of xlsx, pdf, docx or zip.")
    prepared = _prepare(claim)
    output = Path(os.path.abspath(out_dir))
    for component in (output, *output.parents):
        if component.is_symlink():
            raise ReportError("Report output directories must not contain symlinks.")
    output.mkdir(parents=True, exist_ok=True)
    requested = set(formats)
    required = requested | {"xlsx", "pdf"} if "zip" in requested else requested
    paths = {item: output / f"claim_{prepared['code']}.{item}" for item in requested}
    for path in paths.values():
        if path.is_symlink() or path.exists():
            raise ReportError(f"Output already exists; choose a new directory: {path.name}")
    with tempfile.TemporaryDirectory(prefix=".claim-report-", dir=output) as temporary:
        staging = Path(temporary)
        generated: dict[str, Path] = {}
        for extension, writer in [
            ("xlsx", _write_xlsx),
            ("pdf", _write_pdf),
            ("docx", _write_docx),
        ]:
            if extension in required:
                target = staging / f"claim_{prepared['code']}.{extension}"
                writer(prepared, target)
                generated[extension] = target
        if "zip" in requested:
            target = staging / f"claim_{prepared['code']}.zip"
            with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                for document in generated.values():
                    archive.write(document, document.name)
                for receipt in prepared["receipts"]:
                    archive.writestr(f"receipts/{receipt.name}", receipt.data)
                manifest = {
                    "schema_version": 1,
                    "claim_code": prepared["code"],
                    "currency": prepared["currency"],
                    "policy_sha": prepared["policy_sha"],
                    "totals": prepared["totals"],
                    "receipts": [receipt.manifest() for receipt in prepared["receipts"]],
                    "reports": [
                        {
                            "filename": file.name,
                            "sha256": hashlib.sha256(file.read_bytes()).hexdigest(),
                        }
                        for file in generated.values()
                    ],
                }
                archive.writestr(
                    "manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2)
                )
                if prepared["policy_text"]:
                    archive.writestr("policy_snapshot.yaml", prepared["policy_text"])
            generated["zip"] = target
        for extension, destination in paths.items():
            try:
                # Same-filesystem hard linking publishes atomically without overwriting.
                os.link(generated[extension], destination, follow_symlinks=False)
            except FileExistsError as exc:
                raise ReportError(f"Output already exists: {destination.name}") from exc
    return {extension: str(paths[extension]) for extension in sorted(paths)}
